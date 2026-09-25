import argparse
import json
import os
import re
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import boto3
from botocore.exceptions import BotoCoreError, ClientError
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DEFAULT_PROFILE = os.getenv("CLAIMS_AWS_PROFILE")
DEFAULT_TABLE = os.getenv(
    "CLAIMS_AWS_CLAIMS_TABLE",
    "example-claims-assessments",
)
DEFAULT_DYNAMODB_REGION = os.getenv("CLAIMS_AWS_REGION", "us-east-1")
DEFAULT_EVIDENCE_REGION = os.getenv(
    "CLAIMS_AWS_EVIDENCE_REGION",
    DEFAULT_DYNAMODB_REGION,
)

OUTPUT_DIRECTORY = Path("outputs") / "aws_claim_reports"

NAVY = "17324D"
BLUE = "1F73BE"
LIGHT_BLUE = "EAF3FB"
GREEN = "DFF2E6"
GREEN_TEXT = "167A45"
YELLOW = "FFF4D6"
YELLOW_TEXT = "9A6700"
RED = "FDE5E5"
RED_TEXT = "B42318"
LIGHT_GREY = "F4F6F8"
BORDER_COLOUR = "D0D7DE"
WHITE = "FFFFFF"

THIN_BORDER = Border(
    left=Side(style="thin", color=BORDER_COLOUR),
    right=Side(style="thin", color=BORDER_COLOUR),
    top=Side(style="thin", color=BORDER_COLOUR),
    bottom=Side(style="thin", color=BORDER_COLOUR),
)


def json_safe(value):
    if isinstance(value, Decimal):
        if value % 1 == 0:
            return int(value)
        return float(value)

    if isinstance(value, dict):
        return {
            key: json_safe(item_value)
            for key, item_value in value.items()
        }

    if isinstance(value, list):
        return [json_safe(item) for item in value]

    return value


def display_name(value):
    if value is None:
        return ""

    text = str(value).replace("_", " ").strip()
    return " ".join(word.capitalize() for word in text.split())


def display_value(value):
    value = json_safe(value)

    if value is None:
        return "Not available"

    if isinstance(value, bool):
        return "Yes" if value else "No"

    if isinstance(value, list):
        if not value:
            return "None"
        return ", ".join(str(item) for item in value)

    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)

    return str(value)


def percentage(value):
    if value is None or value == "":
        return "Not available"

    try:
        return f"{float(value):.2f}%"
    except (TypeError, ValueError):
        return display_value(value)


def parse_s3_uri(uri):
    match = re.fullmatch(r"s3://([^/]+)/(.+)", str(uri).strip())

    if not match:
        raise ValueError(f"Invalid S3 URI: {uri}")

    return match.group(1), match.group(2)


def read_s3_json(s3_client, uri):
    bucket, key = parse_s3_uri(uri)
    response = s3_client.get_object(Bucket=bucket, Key=key)
    content = response["Body"].read().decode("utf-8")
    return json.loads(content)


def configure_sheet(sheet, freeze_panes=None):
    sheet.sheet_view.showGridLines = False

    if freeze_panes:
        sheet.freeze_panes = freeze_panes


def add_title(sheet, title, subtitle, maximum_column):
    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=maximum_column,
    )

    title_cell = sheet.cell(row=1, column=1, value=title)
    title_cell.font = Font(
        name="Aptos Display",
        size=20,
        bold=True,
        color=WHITE,
    )
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )
    sheet.row_dimensions[1].height = 34

    sheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=maximum_column,
    )

    subtitle_cell = sheet.cell(row=2, column=1, value=subtitle)
    subtitle_cell.font = Font(
        name="Aptos",
        size=10,
        color="5D6B78",
        italic=True,
    )
    subtitle_cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )
    sheet.row_dimensions[2].height = 24


def style_header_row(sheet, row_number, start_column, end_column):
    for column in range(start_column, end_column + 1):
        cell = sheet.cell(row=row_number, column=column)
        cell.font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=WHITE,
        )
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
        )
        cell.border = THIN_BORDER

    sheet.row_dimensions[row_number].height = 26


def style_table_cell(cell):
    cell.font = Font(name="Aptos", size=10, color=NAVY)
    cell.alignment = Alignment(
        horizontal="left",
        vertical="top",
        wrap_text=True,
    )
    cell.border = THIN_BORDER


def apply_status_style(cell, status):
    normalized = str(status or "").strip().upper()

    if normalized in {
        "PASS",
        "VALID",
        "ELIGIBLE",
        "CONSISTENT",
        "COMPLETED",
        "READY_FOR_AGENT_REVIEW",
    }:
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=GREEN_TEXT,
        )
    elif normalized in {
        "REVIEW",
        "NEEDS_REVIEW",
        "IN_PROGRESS",
        "PENDING",
        "ESCALATE_FOR_REVIEW",
    }:
        cell.fill = PatternFill("solid", fgColor=YELLOW)
        cell.font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=YELLOW_TEXT,
        )
    elif normalized in {
        "FAIL",
        "FAILED",
        "INVALID",
        "INELIGIBLE",
        "INCONSISTENT",
        "REJECTED",
    }:
        cell.fill = PatternFill("solid", fgColor=RED)
        cell.font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=RED_TEXT,
        )


def add_key_value_rows(sheet, start_row, values):
    row = start_row

    for label, value in values:
        label_cell = sheet.cell(row=row, column=1, value=label)
        value_cell = sheet.cell(row=row, column=2, value=display_value(value))

        label_cell.font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=NAVY,
        )
        label_cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        label_cell.border = THIN_BORDER
        label_cell.alignment = Alignment(
            horizontal="left",
            vertical="top",
            wrap_text=True,
        )

        value_cell.font = Font(name="Aptos", size=10, color=NAVY)
        value_cell.border = THIN_BORDER
        value_cell.alignment = Alignment(
            horizontal="left",
            vertical="top",
            wrap_text=True,
        )

        if "status" in label.lower() or "action" in label.lower():
            apply_status_style(value_cell, value)

        row += 1

    return row


def find_document_results(document_result):
    for key in (
        "documents",
        "document_results",
        "validation_results",
        "results",
    ):
        value = document_result.get(key)

        if isinstance(value, dict):
            return value

    possible_documents = {}

    for document_type in ("ktp", "sim", "stnk"):
        value = document_result.get(document_type)

        if isinstance(value, dict):
            possible_documents[document_type] = value

    return possible_documents


def find_field_results(document):
    for key in ("field_results", "fields", "checks"):
        value = document.get(key)

        if isinstance(value, list):
            return value

        if isinstance(value, dict):
            output = []

            for field_name, field_value in value.items():
                if isinstance(field_value, dict):
                    item = dict(field_value)
                    item.setdefault("field_name", field_name)
                    output.append(item)

            return output

    return []


def create_summary_sheet(workbook, claim):
    sheet = workbook.active
    sheet.title = "Claim Summary"
    configure_sheet(sheet)

    add_title(
        sheet,
        "Insurance Claim Assessment Report",
        "Document and vehicle-image assessment prepared for human review",
        2,
    )

    summary_values = [
        ("Claim ID", claim.get("claim_id")),
        ("Policy number", claim.get("policy_number")),
        ("Customer ID", claim.get("customer_id")),
        ("Incident date", claim.get("incident_date")),
        ("Document package status", claim.get("document_package_status")),
        ("Document result ready", claim.get("document_result_ready")),
        ("Image evidence status", claim.get("image_evidence_status")),
        ("Image result ready", claim.get("image_result_ready")),
        (
            "Cross-image consistency",
            claim.get("vehicle_consistency_status"),
        ),
        (
            "Cross-image consistency score",
            percentage(claim.get("vehicle_consistency_score")),
        ),
        (
            "Visible damage severity",
            claim.get("visible_damage_severity"),
        ),
        ("Processing stage", claim.get("processing_stage")),
        ("Next action", claim.get("next_action")),
        (
            "Automated claim decision",
            claim.get("automated_claim_decision", False),
        ),
        ("Last updated", claim.get("updated_at")),
    ]

    next_row = add_key_value_rows(sheet, 4, summary_values)

    notice_cell = sheet.cell(
        row=next_row + 1,
        column=1,
        value="Review notice",
    )
    notice_cell.font = Font(
        name="Aptos",
        size=10,
        bold=True,
        color=YELLOW_TEXT,
    )
    notice_cell.fill = PatternFill("solid", fgColor=YELLOW)
    notice_cell.border = THIN_BORDER

    message_cell = sheet.cell(
        row=next_row + 1,
        column=2,
        value=(
            "This report supports claims-agent review. It does not make "
            "an automated insurance approval or rejection decision."
        ),
    )
    message_cell.font = Font(name="Aptos", size=10, color=NAVY)
    message_cell.fill = PatternFill("solid", fgColor=YELLOW)
    message_cell.border = THIN_BORDER
    message_cell.alignment = Alignment(wrap_text=True, vertical="top")

    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 72


def create_document_sheet(workbook, document_result):
    sheet = workbook.create_sheet("Document Assessment")
    configure_sheet(sheet, freeze_panes="A6")

    add_title(
        sheet,
        "Document Assessment",
        "KTP, SIM, and STNK extraction and validation results",
        9,
    )

    package_status = (
        document_result.get("document_package_status")
        or document_result.get("package_status")
        or document_result.get("status")
    )

    next_action = (
        document_result.get("document_next_action")
        or document_result.get("next_action")
    )

    sheet["A4"] = "Package status"
    sheet["B4"] = display_value(package_status)
    sheet["D4"] = "Next action"
    sheet["E4"] = display_value(next_action)

    for coordinate in ("A4", "D4"):
        sheet[coordinate].font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=NAVY,
        )

    for coordinate, status in (
        ("B4", package_status),
        ("E4", next_action),
    ):
        sheet[coordinate].border = THIN_BORDER
        sheet[coordinate].alignment = Alignment(wrap_text=True)
        apply_status_style(sheet[coordinate], status)

    headers = [
        "Document",
        "Field",
        "Extracted result",
        "Expected value",
        "Confidence",
        "Threshold",
        "Match score",
        "Status",
        "Findings",
    ]

    for column, header in enumerate(headers, start=1):
        sheet.cell(row=6, column=column, value=header)

    style_header_row(sheet, 6, 1, len(headers))

    row = 7
    document_results = find_document_results(document_result)

    for document_type, document in document_results.items():
        field_results = find_field_results(document)

        if not field_results:
            field_results = [
                {
                    "field_name": "Document result",
                    "extracted_value": document.get(
                        "document_type",
                        document_type,
                    ),
                    "status": (
                        document.get("document_status")
                        or document.get("status")
                    ),
                    "findings": document.get("findings", []),
                }
            ]

        for field in field_results:
            extracted = (
                field.get("extracted_value")
                if "extracted_value" in field
                else field.get("value")
            )

            expected = (
                field.get("expected_value")
                if "expected_value" in field
                else field.get("expected")
            )

            confidence = (
                field.get("ocr_confidence_score")
                if "ocr_confidence_score" in field
                else field.get("confidence_score")
            )

            threshold = (
                field.get("confidence_threshold")
                if "confidence_threshold" in field
                else field.get("threshold")
            )

            match_score = field.get("match_score")
            status = field.get("status")
            findings = field.get("findings", [])

            values = [
                str(document_type).upper(),
                display_name(
                    field.get("display_name")
                    or field.get("field_name")
                    or field.get("field")
                ),
                display_value(extracted),
                display_value(expected),
                percentage(confidence),
                percentage(threshold),
                percentage(match_score),
                display_value(status),
                display_value(findings),
            ]

            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=column, value=value)
                style_table_cell(cell)

            apply_status_style(sheet.cell(row=row, column=8), status)
            row += 1

    if row == 7:
        sheet.cell(
            row=row,
            column=1,
            value="Detailed document field results were not available.",
        )

    widths = [14, 28, 30, 30, 14, 14, 14, 16, 38]

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.auto_filter.ref = f"A6:I{max(row - 1, 6)}"


def create_image_sheet(workbook, image_result):
    sheet = workbook.create_sheet("Image Assessment")
    configure_sheet(sheet, freeze_panes="A18")

    add_title(
        sheet,
        "Vehicle-Image Assessment",
        "Four-view evidence validation and cross-image vehicle comparison",
        7,
    )

    consistency = image_result.get(
        "cross_image_vehicle_consistency",
        {},
    )
    evidence = image_result.get("evidence_assessment", {})
    damage = image_result.get("damage_assessment", {})

    summary_values = [
        ("Evidence status", evidence.get("evidence_status")),
        (
            "Vehicle consistency status",
            consistency.get("status"),
        ),
        (
            "Vehicle consistency score",
            percentage(consistency.get("consistency_score")),
        ),
        (
            "Consistency pass threshold",
            percentage(consistency.get("pass_threshold")),
        ),
        (
            "Consistency review threshold",
            percentage(consistency.get("review_threshold")),
        ),
        (
            "Highest visible severity",
            damage.get("highest_visible_severity"),
        ),
        (
            "Final recommendation",
            evidence.get("final_recommendation"),
        ),
        (
            "Hard failures",
            evidence.get("hard_failures", []),
        ),
        (
            "Review flags",
            evidence.get("review_flags", []),
        ),
    ]

    add_key_value_rows(sheet, 4, summary_values)

    sheet["A15"] = "Cross-image comparison"
    sheet["A15"].font = Font(
        name="Aptos Display",
        size=14,
        bold=True,
        color=NAVY,
    )

    headers = [
        "Characteristic",
        "Score",
        "Weight",
        "Comparable",
        "Dominant result",
        "Reason",
        "View observations",
    ]

    for column, header in enumerate(headers, start=1):
        sheet.cell(row=17, column=column, value=header)

    style_header_row(sheet, 17, 1, len(headers))

    row = 18

    for component in consistency.get("comparison_components", []):
        observations = []

        for observation in component.get("observations", []):
            observations.append(
                f"{display_name(observation.get('view'))}: "
                f"{display_value(observation.get('value'))} "
                f"({percentage(observation.get('confidence'))})"
            )

        values = [
            component.get("display_name"),
            percentage(component.get("score")),
            percentage(component.get("weight")),
            display_value(component.get("comparable")),
            component.get("dominant_value", "Not available"),
            component.get("reason"),
            "; ".join(observations) if observations else "Not available",
        ]

        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            style_table_cell(cell)

        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="Visible damage summary")
    sheet.cell(row=row, column=1).font = Font(
        name="Aptos Display",
        size=14,
        bold=True,
        color=NAVY,
    )
    row += 2

    damage_values = [
        ("Damage visible", damage.get("damage_visible")),
        (
            "Highest visible severity",
            damage.get("highest_visible_severity"),
        ),
        (
            "Severity confidence",
            percentage(damage.get("severity_confidence")),
        ),
        ("Damage locations", damage.get("damage_locations", [])),
        (
            "Damage classifications",
            damage.get("damage_classifications", []),
        ),
        ("Damaged parts", damage.get("damaged_parts", [])),
        (
            "Repair recommendations",
            damage.get("repair_recommendations", []),
        ),
        ("Limitations", damage.get("limitations", [])),
    ]

    add_key_value_rows(sheet, row, damage_values)

    widths = [28, 14, 14, 14, 24, 48, 70]

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def flatten_json(value, prefix=""):
    rows = []

    if isinstance(value, dict):
        for key, item in value.items():
            path = f"{prefix}.{key}" if prefix else str(key)
            rows.extend(flatten_json(item, path))
    elif isinstance(value, list):
        if not value:
            rows.append((prefix, "[]"))
        else:
            for index, item in enumerate(value):
                path = f"{prefix}[{index}]"
                rows.extend(flatten_json(item, path))
    else:
        rows.append((prefix, display_value(value)))

    return rows


def create_raw_details_sheet(
    workbook,
    claim,
    document_result,
    image_result,
):
    sheet = workbook.create_sheet("Raw Details")
    configure_sheet(sheet, freeze_panes="A4")

    add_title(
        sheet,
        "Assessment Source Details",
        "Flattened source values retained for traceability",
        3,
    )

    headers = ["Source", "Field path", "Value"]

    for column, header in enumerate(headers, start=1):
        sheet.cell(row=4, column=column, value=header)

    style_header_row(sheet, 4, 1, 3)

    row = 5

    sources = [
        ("DynamoDB claim", claim),
        ("Document validation", document_result),
        ("Image validation", image_result),
    ]

    for source_name, payload in sources:
        for path, value in flatten_json(payload):
            values = [source_name, path, str(value)[:32000]]

            for column, item in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=column, value=item)
                style_table_cell(cell)

            row += 1

    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 62
    sheet.column_dimensions["C"].width = 100
    sheet.auto_filter.ref = f"A4:C{max(row - 1, 4)}"


def load_claim(session, table_name, region, claim_id):
    dynamodb = session.resource("dynamodb", region_name=region)
    table = dynamodb.Table(table_name)

    response = table.get_item(
        Key={"claim_id": claim_id},
        ConsistentRead=True,
    )

    claim = response.get("Item")

    if not claim:
        raise ValueError(
            f"Claim {claim_id} was not found in DynamoDB."
        )

    return json_safe(claim)


def build_report(
    claim_id,
    profile,
    table_name,
    dynamodb_region,
    evidence_region,
    output_path=None,
):
    session = boto3.Session(
        profile_name=profile,
        region_name=dynamodb_region,
    )

    claim = load_claim(
        session,
        table_name,
        dynamodb_region,
        claim_id,
    )

    document_uri = claim.get("document_validation_result_s3_uri")
    image_uri = claim.get("image_validation_result_s3_uri")

    if not document_uri:
        raise ValueError(
            "The claim does not have a document-validation result URI."
        )

    if not image_uri:
        raise ValueError(
            "The claim does not have an image-validation result URI."
        )

    s3_client = session.client("s3", region_name=evidence_region)

    document_result = read_s3_json(s3_client, document_uri)
    image_result = read_s3_json(s3_client, image_uri)

    workbook = Workbook()

    create_summary_sheet(workbook, claim)
    create_document_sheet(workbook, document_result)
    create_image_sheet(workbook, image_result)
    create_raw_details_sheet(
        workbook,
        claim,
        document_result,
        image_result,
    )

    if output_path:
        report_path = Path(output_path)
    else:
        OUTPUT_DIRECTORY.mkdir(parents=True, exist_ok=True)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        report_path = OUTPUT_DIRECTORY / (
            f"{claim_id}_claim_assessment_{timestamp}.xlsx"
        )

    report_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(report_path)

    return report_path.resolve()


def main():
    parser = argparse.ArgumentParser(
        description=(
            "Generate an Excel claims-assessment report from AWS."
        )
    )

    parser.add_argument(
        "claim_id",
        nargs="?",
        help="Claim ID to export",
    )
    parser.add_argument(
        "--profile",
        default=DEFAULT_PROFILE,
        help="AWS CLI profile",
    )
    parser.add_argument(
        "--table",
        default=DEFAULT_TABLE,
        help="DynamoDB table name",
    )
    parser.add_argument(
        "--dynamodb-region",
        default=DEFAULT_DYNAMODB_REGION,
        help="DynamoDB region",
    )
    parser.add_argument(
        "--evidence-region",
        default=DEFAULT_EVIDENCE_REGION,
        help="S3 evidence region",
    )
    parser.add_argument(
        "--output",
        help="Optional output XLSX path",
    )

    arguments = parser.parse_args()

    claim_id = arguments.claim_id

    if not claim_id:
        claim_id = input("Enter claim ID: ").strip()

    if not claim_id:
        raise SystemExit("A claim ID is required.")

    try:
        report_path = build_report(
            claim_id=claim_id,
            profile=arguments.profile,
            table_name=arguments.table,
            dynamodb_region=arguments.dynamodb_region,
            evidence_region=arguments.evidence_region,
            output_path=arguments.output,
        )
    except (BotoCoreError, ClientError) as exc:
        raise SystemExit(
            "AWS request failed. If the SSO session expired, run "
            "refresh the configured AWS credentials and try again.\n"
            f"Details: {exc}"
        ) from exc
    except (ValueError, KeyError, json.JSONDecodeError) as exc:
        raise SystemExit(f"Report generation failed: {exc}") from exc

    print("\nClaim assessment report generated successfully.")
    print(f"Output: {report_path}")


if __name__ == "__main__":
    main()
